"""Фикстуры тестов ассистента.

ВНИМАНИЕ: docx в фикстурах синтетические (собираются python-docx по образцу
«регламент с разделами Heading1 и таблицей»). Реальные документы СНАРК исполнителю
не переданы. Когда они появятся, их нужно положить в `fixtures/` и добавить
параметризованный тест парсера на каждом файле.
"""

from __future__ import annotations

import os
import uuid
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from src.core.config import EmbeddingProvider, LLMProvider, RagStorage, settings
from src.modules.assistant import db as assistant_db
from src.modules.assistant.embeddings import FakeEmbedder, reset_embedder

TEST_DB_ENV = "RAG_TEST_DATABASE_URL"


def build_docx(
    sections: list[tuple[str, list[str | list[list[str]]]]], preamble: str = ""
) -> bytes:
    """Собрать docx: [(заголовок Heading1, [абзац | таблица как список строк])]."""
    document = Document()
    if preamble:
        document.add_paragraph(preamble)
    for heading, blocks in sections:
        document.add_heading(heading, level=1)
        for block in blocks:
            if isinstance(block, str):
                document.add_paragraph(block)
            else:
                table = document.add_table(rows=len(block), cols=len(block[0]))
                for row_index, row in enumerate(block):
                    for col_index, value in enumerate(row):
                        table.cell(row_index, col_index).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


ADAPTATION_DOCX_SECTIONS: list[tuple[str, list[str | list[list[str]]]]] = [
    (
        "1. Общие положения",
        ["Регламент описывает порядок адаптации новых сотрудников компании."],
    ),
    (
        "2. Испытательный срок",
        [
            "Испытательный срок для специалистов составляет три месяца.",
            "За две недели до окончания испытательного срока руководитель заполняет лист оценки.",
            [
                ["Категория", "Срок испытания", "Ответственный"],
                ["Специалист", "3 месяца", "Руководитель отдела"],
                ["Руководитель", "6 месяцев", "Директор по персоналу"],
            ],
        ],
    ),
    (
        "3. Наставничество",
        ["Наставник назначается приказом в первый рабочий день нового сотрудника."],
    ),
]

FINANCE_DOCX_SECTIONS: list[tuple[str, list[str | list[list[str]]]]] = [
    (
        "1. Премирование бухгалтерии",
        ["Квартальная премия сотрудников бухгалтерии рассчитывается по итогам закрытия периода."],
    ),
]


@pytest.fixture
def adaptation_docx() -> bytes:
    return build_docx(ADAPTATION_DOCX_SECTIONS, preamble="Утверждено приказом директора.")


@pytest.fixture
def fake_embedder() -> FakeEmbedder:
    return FakeEmbedder()


PUBLIC_DDL = """
CREATE TYPE user_role AS ENUM ('admin', 'hr_manager', 'employee');
CREATE TABLE departments (id uuid PRIMARY KEY DEFAULT gen_random_uuid(), name text NOT NULL);
CREATE TABLE users (
  id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
  email text NOT NULL,
  role user_role NOT NULL DEFAULT 'employee',
  department_id uuid REFERENCES departments(id),
  is_active boolean NOT NULL DEFAULT true
);
CREATE TABLE employee_profiles (
  user_id uuid PRIMARY KEY REFERENCES users(id),
  phone text, position_title text, birth_date date, start_date date,
  inn text, snils text, address text, contract_end_date date
);
CREATE TABLE documents (
  id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
  title text NOT NULL, version text NOT NULL DEFAULT '1.0',
  file_name text NOT NULL, content_type text NOT NULL, file_path text,
  access text NOT NULL DEFAULT 'public', department_id text,
  rag_status text NOT NULL DEFAULT 'draft'
);
CREATE TABLE knowledge_articles (
  id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
  title text NOT NULL, content text NOT NULL,
  is_published boolean NOT NULL DEFAULT false,
  updated_at timestamptz NOT NULL DEFAULT now()
);
"""


@pytest.fixture
def rag_db(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):  # type: ignore[no-untyped-def]
    """Чистая тестовая БД: минимальные таблицы public + схема rag из миграции Alembic.

    Нужна переменная RAG_TEST_DATABASE_URL (postgresql+asyncpg://… на ПУСТУЮ базу
    с расширением vector, пользователь — владелец базы). Без неё тесты пропускаются.
    """
    url = os.environ.get(TEST_DB_ENV)
    if not url:
        pytest.skip(f"{TEST_DB_ENV} не задан")

    import psycopg
    from alembic.config import Config

    from alembic import command

    sync_url = url.replace("postgresql+asyncpg://", "postgresql://")
    with psycopg.connect(sync_url, autocommit=True) as connection:
        connection.execute("DROP SCHEMA IF EXISTS rag CASCADE")
        connection.execute("DROP SCHEMA public CASCADE")
        connection.execute("CREATE SCHEMA public")
        connection.execute("CREATE EXTENSION IF NOT EXISTS vector")
        connection.execute(PUBLIC_DDL)

    storage_dir = tmp_path / "storage"
    storage_dir.mkdir()
    monkeypatch.setattr(settings, "rag_database_url", url)
    monkeypatch.setattr(settings, "embedding_provider", EmbeddingProvider.FAKE)
    monkeypatch.setattr(settings, "llm_provider", LLMProvider.FAKE)
    monkeypatch.setattr(settings, "rag_storage", RagStorage.LOCAL)
    monkeypatch.setattr(settings, "rag_local_storage_dir", storage_dir)
    monkeypatch.setattr(settings, "rag_min_score", 0.2)
    assistant_db.get_rag_engine.cache_clear()
    reset_embedder()

    root = Path(__file__).resolve().parents[4]
    command.upgrade(Config(str(root / "alembic_rag.ini")), "head")

    class Database:
        def __init__(self) -> None:
            self.storage_dir = storage_dir

        def execute(self, sql: str, params: tuple | None = None) -> list[tuple]:
            with psycopg.connect(sync_url, autocommit=True) as connection:
                cursor = connection.execute(sql, params)
                return cursor.fetchall() if cursor.description else []

        def add_department(self, name: str) -> uuid.UUID:
            return self.execute("INSERT INTO departments (name) VALUES (%s) RETURNING id", (name,))[
                0
            ][0]

        def add_user(
            self, role: str = "employee", department_id: uuid.UUID | None = None, **profile: object
        ) -> uuid.UUID:
            user_id = self.execute(
                "INSERT INTO users (email, role, department_id) VALUES (%s, %s, %s) RETURNING id",
                (f"{uuid.uuid4()}@snark.test", role, department_id),
            )[0][0]
            if profile:
                columns = ", ".join(profile)
                marks = ", ".join(["%s"] * len(profile))
                self.execute(
                    f"INSERT INTO employee_profiles (user_id, {columns}) VALUES (%s, {marks})",
                    (user_id, *profile.values()),
                )
            return user_id

        def add_document(
            self,
            title: str,
            data: bytes,
            rag_status: str = "actual",
            access: str = "public",
            department_id: uuid.UUID | None = None,
            version: str = "1.0",
        ) -> uuid.UUID:
            key = f"documents/{uuid.uuid4()}.docx"
            path = storage_dir / key
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
            return self.execute(
                "INSERT INTO documents (title, version, file_name, content_type, file_path, "
                "access, department_id, rag_status) VALUES (%s, %s, %s, %s, %s, %s, %s, %s) "
                "RETURNING id",
                (
                    title,
                    version,
                    "file.docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key,
                    access,
                    str(department_id) if department_id else None,
                    rag_status,
                ),
            )[0][0]

    yield Database()

    assistant_db.get_rag_engine.cache_clear()
    reset_embedder()
