"""Синтетический регламент для E2E и локальной проверки ассистента.

    python scripts/make_demo_docx.py <каталог>

Это НЕ документ СНАРК: текст выдуман, чтобы проверить конвейер
(Heading1-разделы, таблица, цитирование раздела). Реальные документы
загружаются через `pnpm seed:content`.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

TITLE = "E2E Регламент адаптации (демо)"


def build(target: Path) -> Path:
    document = Document()
    document.add_paragraph(
        "Демонстрационный документ для проверки ассистента. Не является регламентом СНАРК."
    )

    document.add_heading("1. Общие положения", level=1)
    document.add_paragraph(
        "Регламент описывает порядок адаптации новых сотрудников в первые месяцы работы."
    )

    document.add_heading("2. Испытательный срок", level=1)
    document.add_paragraph("Испытательный срок для специалистов составляет три месяца.")
    document.add_paragraph(
        "За две недели до окончания испытательного срока руководитель отдела заполняет лист оценки "
        "и передаёт его в отдел кадров."
    )
    table = document.add_table(rows=3, cols=3)
    for row_index, row in enumerate(
        [
            ["Категория", "Срок испытания", "Ответственный"],
            ["Специалист", "3 месяца", "Руководитель отдела"],
            ["Руководитель", "6 месяцев", "Директор по персоналу"],
        ]
    ):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value

    document.add_heading("3. Наставничество", level=1)
    document.add_paragraph(
        "Наставник назначается приказом в первый рабочий день нового сотрудника."
    )
    document.add_paragraph(
        "Наставник проводит встречи с новым сотрудником не реже одного раза в неделю."
    )

    target.mkdir(parents=True, exist_ok=True)
    path = target / f"{TITLE}.docx"
    document.save(path)
    return path


if __name__ == "__main__":
    print(build(Path(sys.argv[1] if len(sys.argv) > 1 else ".")))
